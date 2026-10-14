import os
import sqlite3
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

try:
  from plyer import share
except ImportError:
  share = None


def compartilhar_arquivo(caminho_arquivo):
  caminho_abs = os.path.abspath(caminho_arquivo)
  if not os.path.exists(caminho_abs):
    return
  shared_success = False
  if share is not None:
    try:
      if hasattr(share, "share"):
        share.share(filepath=caminho_abs)
      elif hasattr(share, "share_file"):
        share.share_file(filepath=caminho_abs)
      shared_success = True
    except Exception as e:
      print(f"Erro Plyer: {e}")

  if not shared_success:
    try:
      import platform
      import subprocess

      sistema = platform.system()
      if sistema == "Windows":
        os.startfile(caminho_abs)
      elif sistema == "Darwin":
        subprocess.run(["open", caminho_abs])
      else:
        subprocess.run(["xdg-open", caminho_abs])
    except Exception as err:
      print(f"Erro ao abrir arquivo localmente: {err}")


def exportar_para_excel(
    db_name="frequencia_algo_a_mais.db",
    output_file="relatorio_frequencia.xlsx",
    compartilhar=True,
):
  conn = sqlite3.connect(db_name)
  cursor = conn.cursor()
  cursor.execute(
      "SELECT data, local, membros, total_participantes FROM reunioes ORDER BY"
      " id ASC"
  )
  dados = cursor.fetchall()
  conn.close()

  if not dados:
    return False

  wb = openpyxl.Workbook()
  ws_resumo = wb.active
  ws_resumo.title = "Resumo & Dashboard"
  ws_dados = wb.create_sheet(title="Histórico de Reuniões")

  font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  fill_header = PatternFill(
      start_color="1F2937", end_color="1F2937", fill_type="solid"
  )
  thin_border = Side(style="thin", color="D1D5DB")
  border_all = Border(
      left=thin_border, right=thin_border, top=thin_border, bottom=thin_border
  )

  headers = ["Data", "Local", "Membros", "Total Participantes"]
  ws_dados.append(headers)

  for col_num in range(1, len(headers) + 1):
    cell = ws_dados.cell(row=1, column=col_num)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center")

  for row in dados:
    ws_dados.append(row)

  num_rows = len(dados) + 1
  for r in range(2, num_rows + 1):
    for c in range(1, 5):
      cell = ws_dados.cell(row=r, column=c)
      cell.border = border_all
      cell.alignment = (
          Alignment(horizontal="right")
          if c > 2
          else Alignment(horizontal="center")
      )

  ws_resumo["B2"] = "Métrica"
  ws_resumo["C2"] = "Valor"
  ws_resumo["B2"].font = font_header
  ws_resumo["B2"].fill = fill_header
  ws_resumo["C2"].font = font_header
  ws_resumo["C2"].fill = fill_header

  metricas = [
      ("Reuniões Realizadas", f"=COUNTA('Histórico de Reuniões'!A2:A{num_rows})"),
      (
          "Média de Presença Total",
          f"=ROUND(AVERAGE('Histórico de Reuniões'!D2:D{num_rows}), 0)",
      ),
  ]

  for idx, (m_label, m_formula) in enumerate(metricas, start=3):
    ws_resumo[f"B{idx}"] = m_label
    ws_resumo[f"C{idx}"] = m_formula
    ws_resumo[f"B{idx}"].border = border_all
    ws_resumo[f"C{idx}"].border = border_all
    ws_resumo[f"C{idx}"].alignment = Alignment(horizontal="center")

  chart = BarChart()
  chart.type = "col"
  chart.style = 10
  chart.title = "Evolução de Participantes por Reunião"
  chart.y_axis.title = "Total Participantes"
  chart.x_axis.title = "Data / Reunião"

  data_ref = Reference(ws_dados, min_col=4, min_row=1, max_row=num_rows)
  cats_ref = Reference(ws_dados, min_col=1, min_row=2, max_row=num_rows)
  chart.add_data(data_ref, titles_from_data=True)
  chart.set_categories(cats_ref)
  ws_resumo.add_chart(chart, "E2")

  for ws in [ws_resumo, ws_dados]:
    for col in ws.columns:
      max_len = max(len(str(cell.value or "")) for cell in col)
      col_letter = openpyxl.utils.get_column_letter(col[0].column)
      ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  caminho_completo = os.path.abspath(output_file)
  wb.save(caminho_completo)

  if compartilhar:
    compartilhar_arquivo(caminho_completo)

  return True


def criar_grafico_torres_pdf(reunioes):
    drawing = Drawing(520, 150)
    drawing.add(
        Rect(
            0,
            0,
            520,
            150,
            fillColor=colors.HexColor("#111827"),
            strokeColor=None,
            rx=8,
            ry=8,
        )
    )

    drawing.add(
        String(
            20,
            125,
            "Distribuição de Presença por Reunião",
            fontName="Helvetica-Bold",
            fontSize=11,
            fillColor=colors.white,
        )
    )

    chart = VerticalBarChart()
    chart.x = 30
    chart.y = 25
    chart.height = 85
    chart.width = 460
    chart.data = [[r[5] for r in reunioes]]

    # Remove a data e exibe apenas o nome limpo do local nas legendas do gráfico
    chart.categoryAxis.categoryNames = [str(r[1]).strip().title() for r in reunioes]

    # Ajusta o tamanho da fonte e inclinação para evitar sobreposição
    chart.categoryAxis.labels.fontSize = 7
    chart.categoryAxis.labels.fillColor = colors.HexColor("#9CA3AF")
    chart.categoryAxis.labels.dy = -12
    chart.categoryAxis.labels.angle = 20  # Inclina levemente o texto para caber perfeitamente

    chart.valueAxis.valueMin = 0
    chart.valueAxis.labels.fontSize = 8
    chart.valueAxis.labels.fillColor = colors.HexColor("#9CA3AF")
    chart.bars[0].fillColor = colors.HexColor("#06B6D4")

    drawing.add(chart)
    return drawing


def exportar_para_pdf(
    db_name="frequencia_algo_a_mais.db",
    output_file="relatorio_frequencia.pdf",
    compartilhar=True,
):
  conn = sqlite3.connect(db_name)
  cursor = conn.cursor()

  cursor.execute("""
        SELECT 
            COUNT(*), 
            COALESCE(ROUND(AVG(total_participantes)), 0), 
            COALESCE(ROUND(AVG(membros)), 0)
        FROM reunioes
    """)
  total_reunioes, media_presenca, media_membros = cursor.fetchone()

  cursor.execute(
      "SELECT data, local, membros, convidados_novos, convidados_total,"
      " total_participantes FROM reunioes ORDER BY id DESC"
  )
  reunioes = cursor.fetchall()
  conn.close()

  if not reunioes:
    return False

  caminho_completo = os.path.abspath(output_file)
  doc = SimpleDocTemplate(
      caminho_completo,
      pagesize=A4,
      rightMargin=36,
      leftMargin=36,
      topMargin=36,
      bottomMargin=36,
  )
  elements = []
  styles = getSampleStyleSheet()

  title_style = ParagraphStyle(
      "DocTitle",
      parent=styles["Normal"],
      fontName="Helvetica-Bold",
      fontSize=20,
      leading=24,
      textColor=colors.HexColor("#1F2937"),
  )
  subtitle_style = ParagraphStyle(
      "DocSubTitle",
      parent=styles["Normal"],
      fontName="Helvetica",
      fontSize=10,
      leading=12,
      textColor=colors.HexColor("#6B7280"),
  )
  section_style = ParagraphStyle(
      "SectionTitle",
      parent=styles["Normal"],
      fontName="Helvetica-Bold",
      fontSize=13,
      leading=16,
      textColor=colors.HexColor("#111827"),
      spaceAfter=8,
  )
  card_title = ParagraphStyle(
      "CardTitle",
      parent=styles["Normal"],
      fontName="Helvetica",
      fontSize=9,
      leading=11,
      textColor=colors.HexColor("#4B5563"),
      alignment=1,
  )
  card_val = ParagraphStyle(
      "CardVal",
      parent=styles["Normal"],
      fontName="Helvetica-Bold",
      fontSize=16,
      leading=18,
      textColor=colors.HexColor("#111827"),
      alignment=1,
  )

  elements.append(
      Paragraph("Relatório de Frequência - ALGO A MAIS", title_style)
  )
  elements.append(
      Paragraph("Resumo Executivo, Gráficos e Histórico", subtitle_style)
  )
  elements.append(Spacer(1, 12))
  elements.append(
      HRFlowable(
          width="100%",
          thickness=1,
          color=colors.HexColor("#E5E7EB"),
          spaceAfter=14,
      )
  )

  elements.append(Paragraph("Resumo das Métricas (Médias)", section_style))

  cards_data = [
      [
          Paragraph("Reuniões Realizadas", card_title),
          Paragraph("Média Presença Total", card_title),
          Paragraph("Média de Membros", card_title),
      ],
      [
          Paragraph(str(int(total_reunioes)), card_val),
          Paragraph(str(int(media_presenca)), card_val),
          Paragraph(str(int(media_membros)), card_val),
      ],
  ]

  cards_table = Table(cards_data, colWidths=[173, 173, 173])
  cards_table.setStyle(
      TableStyle([
          ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F9FAFB")),
          ("ALIGN", (0, 0), (-1, -1), "CENTER"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#E5E7EB")),
          ("INNERGRID", (0, 0), (-1, -1), 1, colors.HexColor("#E5E7EB")),
          ("TOPPADDING", (0, 0), (-1, -1), 8),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
      ])
  )
  elements.append(cards_table)
  elements.append(Spacer(1, 16))

  elements.append(Paragraph("Gráfico de Desempenho por Reunião", section_style))
  elements.append(criar_grafico_torres_pdf(list(reversed(reunioes))))
  elements.append(Spacer(1, 16))

  elements.append(Paragraph("Histórico de Reuniões", section_style))

  table_data = [["Data", "Local", "Membros", "Total Reunião"]]
  for reg in reunioes:
    table_data.append([str(reg[0]), str(reg[1]).strip().title(), str(reg[2]), str(reg[5])])

  hist_table = Table(table_data, colWidths=[110, 150, 110, 150])
  hist_table.setStyle(
      TableStyle([
          ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
          ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
          ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
          ("FONTSIZE", (0, 0), (-1, 0), 9),
          ("ALIGN", (0, 0), (-1, -1), "CENTER"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          (
              "ROWBACKGROUNDS",
              (0, 1),
              (-1, -1),
              [colors.white, colors.HexColor("#F9FAFB")],
          ),
          ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
          ("TOPPADDING", (0, 0), (-1, -1), 6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
      ])
  )
  elements.append(hist_table)

  doc.build(elements)

  if compartilhar:
    compartilhar_arquivo(caminho_completo)

  return True